from lectureweaver import materials as M
from lectureweaver.search import search
from conftest import make_course, make_lecture


def test_import_docx_pptx_code(kb, tmp_path):
    import docx
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("实验三 文件操作", 1)
    d.add_paragraph("要求：用 FileInputStream 实现文件拷贝，并比较缓冲区大小对耗时的影响。")
    d.save(tmp_path / "lab.docx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "BufferedReader"
    s.placeholders[1].text = "readLine() 一次读一行"
    prs.save(tmp_path / "ch10.pptx")
    (tmp_path / "Copy.java").write_text("class Copy { void run() { fos.write(data, 0, length); } }", encoding="utf-8")
    out = [M.import_file(kb, tmp_path / f) for f in ("lab.docx", "ch10.pptx", "Copy.java")]
    assert [o[0]["kind"] for o in out] == ["document", "slides", "code"]
    again, new, _ = M.import_file(kb, tmp_path / "lab.docx")
    assert not new and again["id"] == out[0][0]["id"]
    assert "slide 1" in M.material_text(kb, out[1][0]["id"], "1")


def test_link_suggest_and_search(kb, tmp_path):
    cid = make_course(kb)
    lid = make_lecture(kb, cid)
    make_lecture(kb, cid, lid="2026-10-13_1234567_p3-4", date="2026-10-13")
    kb.index_replace("transcript", lid, [("[10:10:00] read 方法返回成功读取的字节个数", lid, cid, None)])
    f = tmp_path / "guide.md"
    f.write_text("实验：read 方法 成功读取 字节个数", encoding="utf-8")
    m, _, _ = M.import_file(kb, f)
    M.link(kb, m["id"], "course", cid, "lab_guide")
    sug = M.suggest_lectures(kb, m["id"], cid)
    assert sug and sug[0]["lecture"] in (lid, "2026-10-13_1234567_p3-4")
    hits = search(kb, "字节个数")
    kinds = {h["kind"] for h in hits}
    assert {"transcript", "material"} <= kinds
    t = next(h for h in hits if h["kind"] == "transcript")
    assert t["clock"] == "10:10:00" and t["date"] == "2026-10-20"
    assert search(kb, "读取")  # 2-character query uses LIKE


def test_garbled_text_layer_is_detected():
    # what pypdf returns for Beamer slides whose CJK font has a private encoding
    assert M.is_garbled("V˙ØŒnÚO\nEmail: teacher@example.edu.cn\nn˘ E339\n2026c ¢\n(n˘) V˙ØŒnÚO 2026c ¢ 1 / 789")
    assert M.is_garbled("Œ‰˘¥%4‰n\n8„\n1 ¯ﬂ9ÙV˙\n2 ¯Cþ'Ù9ÙŒiA")
    assert not M.is_garbled("设随机变量 X 的数学期望 E(X)=μ，方差 σ²，则对任意 ε>0 有 P{|X−μ|≥ε}≤σ²/ε²。")
    assert not M.is_garbled("Le café est très élégant, où êtes-vous? Noël à Paris. Die Größe über alles.")
    assert not M.is_garbled("public static void main(String[] args) { System.out.println(\"hi\"); }")


class _FakePage:
    def __init__(self, text):
        self.text = text

    def extract_text(self):
        return self.text


def test_garbled_pdf_is_ocred_and_refresh_reextracts(kb, tmp_path, monkeypatch):
    import pypdf

    pages = ["Œ‰˘¥%4‰n 8„ 1 ¯ﬂ9ÙV˙ 2 ¯Cþ'Ù9ÙŒiA"] * 3
    monkeypatch.setattr(pypdf, "PdfReader", lambda p: type("R", (), {"pages": [_FakePage(t) for t in pages]})())
    ocr_calls = []

    def fake_ocr(path, units, why="扫描版 PDF"):
        ocr_calls.append(why)
        return [(f"p.{i}", f"定理 4.{i} 切比雪夫不等式") for i in range(1, len(units) + 1)], f"{why}，文字来自本地 OCR（可能有识别错误）"

    monkeypatch.setattr(M, "_ocr_pdf", fake_ocr)
    f = tmp_path / "ch4.pdf"
    f.write_bytes(b"%PDF-1.4 fake")
    m, new, warn = M.import_file(kb, f)
    assert new and "乱码" in warn and ocr_calls == ["PDF 文字层是乱码（字体自定义编码）"]
    assert "切比雪夫不等式" in M.material_text(kb, m["id"], "2")

    # a readable text layer is kept as is, unless OCR is forced
    pages[:] = ["设随机变量 X 的数学期望与方差存在，则对任意 ε>0 有切比雪夫不等式。"] * 2
    m2, warn = M.refresh(kb, m["id"])
    assert m2["units"] == 2 and not warn and "数学期望与方差" in M.material_text(kb, m["id"])
    assert search(kb, "数学期望与方差")
    _, warn = M.refresh(kb, m["id"], force_ocr=True)
    assert "按要求 OCR" in warn
