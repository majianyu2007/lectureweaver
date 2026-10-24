"""User-provided course materials: PPT/PDF/Word/Markdown/text/code.

Each file is copied into the knowledge base (content-addressed, so importing
the same file twice is harmless), its text is extracted with page/slide
markers so notes can cite "p.3", and it is linked to a course and/or lectures
with a role. Material text is always labelled as material, never as what the
teacher said in class.
"""

from __future__ import annotations

import hashlib
import math
import re
import shutil
import time
from collections import Counter
from pathlib import Path

from .kb import KB, dumps

CODE_EXT = {".java": "java", ".py": "python", ".c": "c", ".h": "c", ".cpp": "cpp", ".hpp": "cpp", ".js": "javascript", ".ts": "typescript",
            ".sql": "sql", ".go": "go", ".rs": "rust", ".kt": "kotlin", ".m": "matlab", ".r": "r", ".sh": "bash", ".html": "html", ".css": "css",
            ".json": "json", ".xml": "xml", ".yaml": "yaml", ".yml": "yaml", ".ipynb": "json"}
ROLES = ["textbook", "lab_guide", "slides", "code", "reference", "homework", "exam", "notes"]


class MaterialError(Exception):
    pass


def sha256_file(p: Path) -> str:
    h = hashlib.sha256()
    with open(p, "rb") as f:
        for b in iter(lambda: f.read(1 << 20), b""):
            h.update(b)
    return h.hexdigest()


def extract(path: Path) -> tuple[str, list[tuple[str, str]], str]:
    """Return (kind, [(unit label, text)], warning)."""
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader

        rd = PdfReader(str(path))
        units = [(f"p.{i}", (pg.extract_text() or "").strip()) for i, pg in enumerate(rd.pages, 1)]
        empty = sum(1 for _, t in units if len(t) < 20)
        warn = ""
        if units and empty / len(units) > 0.6:
            units, warn = _ocr_pdf(path, units)
        return "pdf", units, warn
    if ext in (".pptx",):
        from pptx import Presentation

        prs = Presentation(str(path))
        units = []
        for i, slide in enumerate(prs.slides, 1):
            parts = []
            for shp in slide.shapes:
                if shp.has_text_frame:
                    parts.append(shp.text_frame.text)
                if getattr(shp, "has_table", False) and shp.has_table:
                    for row in shp.table.rows:
                        parts.append(" | ".join(c.text for c in row.cells))
            if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
                parts.append("[备注] " + slide.notes_slide.notes_text_frame.text)
            units.append((f"slide {i}", "\n".join(p for p in parts if p.strip())))
        return "slides", units, ""
    if ext in (".ppt", ".doc"):
        raise MaterialError(f"{ext} 是旧版 Office 格式，请先另存为 .pptx/.docx/.pdf 再导入")
    if ext == ".docx":
        import docx

        d = docx.Document(str(path))
        out, buf, n = [], [], 1
        for para in d.paragraphs:
            t = para.text.strip()
            if para.style.name.lower().startswith("heading") and buf:
                out.append((f"§{n}", "\n".join(buf)))
                buf, n = [], n + 1
            if t:
                buf.append(("#" * _level(para.style.name) + " " if para.style.name.lower().startswith("heading") else "") + t)
        for ti, table in enumerate(d.tables, 1):
            buf.append(f"[表格{ti}]")
            for row in table.rows:
                buf.append(" | ".join(c.text.strip() for c in row.cells))
        if buf:
            out.append((f"§{n}", "\n".join(buf)))
        return "document", out, ""
    if ext in CODE_EXT:
        return "code", [("file", path.read_text("utf-8", errors="replace"))], ""
    if ext in (".md", ".markdown", ".txt", ".csv", ".tex"):
        text = path.read_text("utf-8", errors="replace")
        return "text", [("file", text)], ""
    raise MaterialError(f"暂不支持 {ext} 文件（支持 PDF/PPTX/DOCX/Markdown/文本/常见代码文件）")


def _level(style: str) -> int:
    m = re.search(r"(\d)", style)
    return min(int(m.group(1)), 4) + 1 if m else 2


def _ocr_pdf(path: Path, units):
    """Scanned PDF: render pages with PDFKit and OCR them locally (no upload)."""
    from .ocr import OcrUnavailable, lines_text, ocr_pdf

    try:
        pages = ocr_pdf(path)
    except (OcrUnavailable, OSError) as e:
        return units, f"扫描版 PDF，未能提取文字：{e}"
    if not pages:
        return units, "扫描版 PDF，未能提取文字"
    return [(f"p.{i}", lines_text(p)) for i, p in enumerate(pages, 1)], "扫描版 PDF，文字来自本地 OCR（可能有识别错误）"


def import_file(kb: KB, src: Path, title: str | None = None, note: str | None = None) -> tuple[dict, bool, str]:
    src = Path(src).expanduser().resolve()
    if not src.is_file():
        raise MaterialError(f"文件不存在：{src}")
    digest = sha256_file(src)
    existing = kb.q1("SELECT * FROM materials WHERE sha256=?", digest)
    if existing:
        return dict(existing), False, ""
    kind, units, warn = extract(src)
    mid = f"m_{digest[:12]}"
    d = kb.abs(f"materials/{mid}")
    d.mkdir(parents=True, exist_ok=True)
    dst = d / src.name
    shutil.copy2(src, dst)
    lines = [f"# {title or src.stem}", "", f"> 资料原文件：{src.name}（导入于 {time.strftime('%Y-%m-%d')}）。以下为提取的文字，图片/公式可能缺失。", ""]
    for label, text in units:
        lines.append(f"## {label}")
        lines.append("")
        if kind == "code":
            lines.append(f"```{CODE_EXT.get(src.suffix.lower(), '')}")
            lines.append(text)
            lines.append("```")
        else:
            lines.append(text or "（本页无可提取文字）")
        lines.append("")
    (d / "content.md").write_text("\n".join(lines), "utf-8")
    kb.x(
        "INSERT INTO materials(id,sha256,filename,kind,title,path,text_path,units,imported_at,note) VALUES(?,?,?,?,?,?,?,?,?,?)",
        mid, digest, src.name, kind, title or src.stem, f"materials/{mid}/{src.name}", f"materials/{mid}/content.md", len(units), time.time(), note,
    )
    kb.db.commit()
    kb.index_replace("material", mid, [(f"[{label}] {text}", None, None, None) for label, text in units])
    return dict(kb.q1("SELECT * FROM materials WHERE id=?", mid)), True, warn


def link(kb: KB, mid: str, target_type: str, target_id: str, role: str | None, note: str | None = None) -> None:
    if target_type not in ("course", "lecture", "concept"):
        raise MaterialError("关联对象只能是 course/lecture/concept")
    if role and role not in ROLES:
        raise MaterialError(f"角色可选：{', '.join(ROLES)}")
    table = {"course": "courses", "lecture": "lectures", "concept": "concepts"}[target_type]
    if not kb.q1(f"SELECT 1 FROM {table} WHERE id=?", target_id):
        raise MaterialError(f"{target_type} {target_id} 不存在")
    kb.x(
        "INSERT INTO material_links(material_id,target_type,target_id,role,note) VALUES(?,?,?,?,?) ON CONFLICT(material_id,target_type,target_id) DO UPDATE SET role=excluded.role, note=excluded.note",
        mid, target_type, target_id, role, note,
    )
    if target_type == "course":
        # make course-level search hits know the course
        kb.x("UPDATE search SET course_id=? WHERE kind='material' AND ref=?", target_id, mid)
    kb.db.commit()


def unlink(kb: KB, mid: str, target_type: str, target_id: str) -> None:
    kb.x("DELETE FROM material_links WHERE material_id=? AND target_type=? AND target_id=?", mid, target_type, target_id)
    kb.db.commit()


_TOKEN = re.compile(r"[A-Za-z_][A-Za-z0-9_]{2,}|[一-鿿]{2,}")


def _terms(text: str) -> Counter:
    c: Counter = Counter()
    for tok in _TOKEN.findall(text):
        if re.match(r"[一-鿿]", tok):
            for i in range(len(tok) - 1):
                c[tok[i : i + 2]] += 1
        else:
            c[tok.lower()] += 1
    return c


def suggest_lectures(kb: KB, mid: str, course_id: str, top: int = 5) -> list[dict]:
    """Rank the course's processed lectures by TF-IDF similarity to the material."""
    import json

    m = kb.q1("SELECT * FROM materials WHERE id=?", mid)
    mtext = kb.abs(m["text_path"]).read_text("utf-8")
    docs = {}
    for l in kb.q("SELECT * FROM lectures WHERE course_id=? ORDER BY date", course_id):
        d = kb.abs(l["dir"])
        parts = []
        try:
            parts += [s["text"] for s in json.loads((d / "transcript.json").read_text("utf-8")).get("segments", [])]
        except (OSError, ValueError):
            pass
        try:
            parts += [f.get("ocr", "") for f in json.loads((d / "frames.json").read_text("utf-8")).get("frames", [])]
        except (OSError, ValueError):
            pass
        if parts:
            docs[l["id"]] = (_terms("\n".join(parts)), l)
    if not docs:
        return []
    n = len(docs)
    df: Counter = Counter()
    for tf, _ in docs.values():
        df.update(set(tf))
    mt = _terms(mtext)

    def vec(tf):
        return {t: (1 + math.log(c)) * math.log((n + 1) / (df.get(t, 0) + 1)) for t, c in tf.items() if df.get(t, 0) < n}

    mv = vec(mt)
    mnorm = math.sqrt(sum(v * v for v in mv.values())) or 1
    scored = []
    for lid, (tf, l) in docs.items():
        lv = vec(tf)
        dot = sum(mv[t] * lv.get(t, 0) for t in mv)
        norm = math.sqrt(sum(v * v for v in lv.values())) or 1
        shared = sorted((t for t in mv if t in lv), key=lambda t: -mv[t] * lv[t])[:8]
        scored.append({"lecture": lid, "date": l["date"], "week": l["week"], "score": round(dot / (mnorm * norm), 3), "shared_terms": shared})
    scored.sort(key=lambda x: -x["score"])
    return scored[:top]


def material_text(kb: KB, mid: str, units: str | None = None) -> str:
    m = kb.q1("SELECT * FROM materials WHERE id=?", mid)
    if not m:
        raise MaterialError(f"资料 {mid} 不存在")
    text = kb.abs(m["text_path"]).read_text("utf-8")
    if not units:
        return text
    wanted = set()
    for part in units.split(","):
        a, _, b = part.partition("-")
        for i in range(int(a), int(b or a) + 1):
            wanted.add(i)
    out, keep = [], True
    for ln in text.splitlines():
        hm = re.match(r"## (?:p\.|slide |§)(\d+)", ln)
        if hm:
            keep = int(hm.group(1)) in wanted
        if keep or ln.startswith("# "):
            out.append(ln)
    return "\n".join(out)


__all__ = ["import_file", "link", "unlink", "suggest_lectures", "material_text", "MaterialError", "ROLES", "dumps"]
